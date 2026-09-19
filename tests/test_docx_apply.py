"""Phase 1 slice 4: apply edit_docx ops with zip-preserving CAS save."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from sqlalchemy.exc import IntegrityError

from agent_console.db.models import User
from agent_console.repositories.files import FileRepository
from agent_console.services.docx_blocks import assign_fresh_manifest, block_texts
from agent_console.services.docx_edit import EditOp, apply_ops, package_entry_digests


def _minimal_docx(*paragraphs: str) -> bytes:
    doc = Document()
    if doc.paragraphs:
        doc.paragraphs[0].text = paragraphs[0] if paragraphs else ""
        rest = paragraphs[1:]
    else:
        rest = paragraphs
    for text in rest:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _paragraph_c14n_list(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    return [ET.canonicalize(ET.tostring(p, encoding="unicode")) for p in root.iter(f"{ns}p")]


def test_apply_only_changes_document_xml_bytes() -> None:
    data = _minimal_docx("Due in 45 days.", "Leave me alone.")
    manifest = assign_fresh_manifest(data, generation=1)
    before = package_entry_digests(data)
    out, new_manifest, diffs = apply_ops(
        data,
        manifest,
        [EditOp(op="replace", block_id="g1:p_0001", old="45 days", new="30 days")],
    )
    after = package_entry_digests(out)
    assert set(before) == set(after)
    changed = [name for name in before if before[name] != after[name]]
    assert changed == ["word/document.xml"]
    assert diffs[0].after == "Due in 30 days."
    assert new_manifest.generation == 1


def test_untouched_paragraphs_are_xml_identical() -> None:
    data = _minimal_docx("Edit me", "Untouched A", "Untouched B")
    manifest = assign_fresh_manifest(data, generation=1)
    before_paras = _paragraph_c14n_list(data)
    out, _, _ = apply_ops(
        data,
        manifest,
        [EditOp(op="replace", block_id="g1:p_0001", old="Edit me", new="Edited")],
    )
    after_paras = _paragraph_c14n_list(out)
    assert len(before_paras) == len(after_paras)
    assert before_paras[1] == after_paras[1]
    assert before_paras[2] == after_paras[2]
    assert before_paras[0] != after_paras[0]


def test_manifest_carry_forward_insert_and_delete() -> None:
    data = _minimal_docx("A", "B", "C")
    manifest = assign_fresh_manifest(data, generation=1)
    assert manifest.next_id == 4
    out, new_manifest, _ = apply_ops(
        data,
        manifest,
        [
            EditOp(op="delete", block_id="g1:p_0003", hash=manifest.blocks[2].content_hash),
            EditOp(
                op="insert",
                relative_to="g1:p_0001",
                position="after",
                content="Inserted",
            ),
        ],
    )
    assert new_manifest.generation == 1
    # Highest id deleted — next_id must not recycle to 3.
    assert new_manifest.next_id == 5
    ids = [b.id for b in new_manifest.blocks]
    assert "g1:p_0003" not in ids
    assert "g1:p_0004" in ids  # insert took next_id 4
    assert block_texts(out) == ["A", "Inserted", "B"]


def test_output_reopens_with_python_docx() -> None:
    data = _minimal_docx("Hello world")
    manifest = assign_fresh_manifest(data, generation=1)
    out, _, _ = apply_ops(
        data,
        manifest,
        [EditOp(op="replace", block_id="g1:p_0001", old="world", new="there")],
    )
    doc = Document(io.BytesIO(out))
    assert "Hello there" in "\n".join(p.text for p in doc.paragraphs)


@pytest.mark.soffice
def test_output_survives_libreoffice_headless(tmp_path: Path) -> None:
    import shutil
    import subprocess

    soffice = shutil.which("soffice")
    if soffice is None:
        pytest.skip("LibreOffice not installed")
    probe = subprocess.run(
        [soffice, "--version"], capture_output=True, text=True, timeout=30
    )
    if probe.returncode != 0:
        pytest.skip(f"LibreOffice stub broken: {probe.stderr.strip() or probe.stdout}")

    data = _minimal_docx("Convert me please")
    manifest = assign_fresh_manifest(data, generation=1)
    out, _, _ = apply_ops(
        data,
        manifest,
        [
            EditOp(
                op="rewrite",
                block_id="g1:p_0001",
                content="Converted",
                hash=manifest.blocks[0].content_hash,
            )
        ],
    )
    src = tmp_path / "in.docx"
    src.write_bytes(out)
    subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_path),
            str(src),
        ],
        check=True,
        capture_output=True,
        timeout=120,
    )
    assert (tmp_path / "in.pdf").is_file()


@pytest.mark.db
async def test_edit_docx_lost_race_leaves_no_orphan_blob(
    files: FileRepository, user: User, blob_dir: Path
) -> None:
    data = _minimal_docx("Race me")
    tip = await files.save(user.id, name="race.docx", data=data)
    await files.save(
        user.id,
        name="race.docx",
        data=data,
        parent_id=tip.id,
    )
    before = {p.name for p in blob_dir.iterdir()}
    # Simulate edit_docx CAS: loser still parents the old tip.
    out, _, _ = apply_ops(
        data,
        assign_fresh_manifest(data, generation=1),
        [EditOp(op="replace", block_id="g1:p_0001", old="Race me", new="Lost")],
    )
    with pytest.raises(IntegrityError):
        await files.save(
            user.id,
            name="race.docx",
            data=out,
            parent_id=tip.id,
        )
    after = {p.name for p in blob_dir.iterdir()}
    assert after == before
