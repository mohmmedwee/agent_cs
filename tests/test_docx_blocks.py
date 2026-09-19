"""Phase 1 slice 1: DOCX block parser + manifest (generations, next_id, hashes)."""

from __future__ import annotations

import io
import os
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest
from docx import Document

from agent_console.repositories.extraction import _docx_text


def _minimal_docx(*paragraphs: str) -> bytes:
    doc = Document()
    # Some python-docx builds start empty; always add explicitly.
    if doc.paragraphs:
        doc.paragraphs[0].text = paragraphs[0] if paragraphs else ""
        rest = paragraphs[1:]
    else:
        rest = paragraphs
    for text in rest:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _legacy_docx_lines(data: bytes) -> list[str]:
    """Frozen copy of extraction._docx_text's per-`w:p` line rules (pre-join).

    Parity tests lock against this so the block parser cannot drift from the
    text the agent already reads today.
    """
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        document = archive.read("word/document.xml")
    root = ElementTree.fromstring(document)
    lines: list[str] = []
    for paragraph in root.iter(f"{_WORD_NS}p"):
        parts: list[str] = []
        for node in paragraph.iter():
            tag = node.tag
            if tag == f"{_WORD_NS}t":
                parts.append(node.text or "")
            elif tag == f"{_WORD_NS}tab":
                parts.append("\t")
            elif tag == f"{_WORD_NS}br":
                parts.append("\n")
        lines.append("".join(parts).strip())
    return lines


def test_block_texts_match_legacy_lines_for_simple_doc() -> None:
    from agent_console.services.docx_blocks import block_texts

    data = _minimal_docx("Hello", "World", "")
    assert block_texts(data) == _legacy_docx_lines(data) == ["Hello", "World", ""]


def test_docx_text_is_join_of_block_texts() -> None:
    from agent_console.services.docx_blocks import block_texts

    data = _minimal_docx("Hello", "", "", "World")
    texts = block_texts(data)
    rebuilt = re.sub(r"\n{3,}", "\n\n", "\n".join(texts)).strip() or None
    assert rebuilt == _docx_text(data)


def test_fresh_manifest_assigns_generation_scoped_ids_and_next_id() -> None:
    from agent_console.services.docx_blocks import (
        Manifest,
        assign_fresh_manifest,
        content_hash,
    )

    data = _minimal_docx("Alpha", "Beta")
    manifest = assign_fresh_manifest(data, generation=1)
    assert isinstance(manifest, Manifest)
    assert manifest.generation == 1
    assert [b.id for b in manifest.blocks] == ["g1:p_0001", "g1:p_0002"]
    assert manifest.next_id == 3
    assert all(b.content_hash == content_hash(b.text) for b in manifest.blocks)
    assert manifest.blocks[0].text == "Alpha"
    assert manifest.blocks[1].text == "Beta"


def test_next_generation_starts_ids_fresh() -> None:
    from agent_console.services.docx_blocks import assign_fresh_manifest

    data = _minimal_docx("Only")
    m2 = assign_fresh_manifest(data, generation=2)
    assert m2.blocks[0].id == "g2:p_0001"
    assert m2.next_id == 2


def test_next_id_survives_delete_of_highest_block() -> None:
    from agent_console.services.docx_blocks import Manifest, assign_fresh_manifest

    data = _minimal_docx("a", "b", "c")
    m = assign_fresh_manifest(data, generation=1)
    carried = Manifest(
        generation=m.generation,
        next_id=m.next_id,
        blocks=m.blocks[:-1],
    )
    assert carried.next_id == 4
    assert [b.id for b in carried.blocks] == ["g1:p_0001", "g1:p_0002"]


def _iter_parity_docx_paths() -> list[Path]:
    raw = os.environ.get("DOCX_PARITY_DIR", "")
    roots = [Path(p) for p in raw.split(os.pathsep) if p.strip()]
    if not roots:
        return []
    out: list[Path] = []
    seen: set[Path] = set()
    for root in roots:
        if not root.is_dir():
            continue
        # Shallow walk only — Desktop trees can be huge; corpus dirs are flat
        # or one-level (e.g. Desktop/DGE). Cap depth at 2 below the root.
        for dirpath, dirnames, filenames in os.walk(root):
            rel = Path(dirpath).relative_to(root)
            depth = 0 if str(rel) == "." else len(rel.parts)
            if depth >= 2:
                dirnames.clear()
            for name in filenames:
                if not name.endswith(".docx") or name.startswith("~$"):
                    continue
                path = Path(dirpath) / name
                if path not in seen:
                    seen.add(path)
                    out.append(path)
    return sorted(out)


def test_block_texts_parity_with_docx_text_on_real_files() -> None:
    """For every corpus DOCX, parser block texts == _docx_text's lines.

    Set DOCX_PARITY_DIR to a colon-separated list of directories outside the
    repo (client documents must not be committed). Skips when unset.
    """
    from agent_console.services.docx_blocks import block_texts

    paths = _iter_parity_docx_paths()
    if not paths:
        pytest.skip("DOCX_PARITY_DIR not set — no external corpus")

    failures: list[str] = []
    checked = 0
    for path in paths:
        try:
            data = path.read_bytes()
        except OSError as exc:
            failures.append(f"{path}: read error {exc}")
            continue
        if not data.startswith(b"PK"):
            continue
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if "word/document.xml" not in zf.namelist():
                    continue
        except zipfile.BadZipFile:
            continue

        checked += 1
        try:
            expected = _legacy_docx_lines(data)
            actual = block_texts(data)
        except Exception as exc:  # noqa: BLE001 — collect all corpus failures
            failures.append(f"{path}: {type(exc).__name__}: {exc}")
            continue
        if actual != expected:
            failures.append(
                f"{path}: mismatch (blocks={len(actual)} lines={len(expected)})"
            )
            continue
        rebuilt = re.sub(r"\n{3,}", "\n\n", "\n".join(actual)).strip() or None
        if rebuilt != _docx_text(data):
            failures.append(f"{path}: join/collapse diverges from _docx_text")

    assert checked > 0, "DOCX_PARITY_DIR set but no readable .docx found"
    assert not failures, (
        f"{len(failures)}/{checked} parity failures:\n" + "\n".join(failures[:30])
    )
