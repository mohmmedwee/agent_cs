"""rewrite allow_format_loss + approval card warning."""

from __future__ import annotations

import io

import pytest
from docx import Document

from agent_console.services.docx_approval import build_approval_card
from agent_console.services.docx_blocks import assign_fresh_manifest
from agent_console.services.docx_edit import (
    FORMAT_LOSS_NOTE,
    EditOp,
    ValidationError,
    apply_ops,
    validate_ops,
)


def _mixed_docx() -> bytes:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("See the ")
    bold = p.add_run("pricing")
    bold.bold = True
    p.add_run(" schedule.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_rewrite_mixed_formatting_requires_allow_format_loss() -> None:
    data = _mixed_docx()
    manifest = assign_fresh_manifest(data, generation=1)
    block = manifest.blocks[0]
    with pytest.raises(ValidationError, match="allow_format_loss"):
        validate_ops(
            manifest,
            [
                EditOp(
                    op="rewrite",
                    block_id=block.id,
                    content="Flat text",
                    hash=block.content_hash,
                )
            ],
            data=data,
        )


def test_rewrite_mixed_with_allow_format_loss_applies_and_cards_warn() -> None:
    data = _mixed_docx()
    manifest = assign_fresh_manifest(data, generation=1)
    block = manifest.blocks[0]
    op = EditOp(
        op="rewrite",
        block_id=block.id,
        content="Flat text",
        hash=block.content_hash,
        allow_format_loss=True,
    )
    result = validate_ops(manifest, [op], data=data)
    card = build_approval_card(
        file_name="x.docx",
        from_version=1,
        diffs=result.diffs,
        ops=[op],
    )
    assert FORMAT_LOSS_NOTE in card["notes"]
    assert FORMAT_LOSS_NOTE in card["changes"][0]["notes"]
    out, _, _ = apply_ops(data, manifest, [op])
    from agent_console.services.docx_blocks import block_texts

    assert block_texts(out)[0] == "Flat text"


def test_rewrite_single_run_does_not_require_flag() -> None:
    doc = Document()
    doc.add_paragraph("Only one run")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    manifest = assign_fresh_manifest(data, generation=1)
    block = manifest.blocks[0]
    validate_ops(
        manifest,
        [
            EditOp(
                op="rewrite",
                block_id=block.id,
                content="Still one",
                hash=block.content_hash,
            )
        ],
        data=data,
    )
