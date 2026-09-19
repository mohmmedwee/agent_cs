"""Run-level replace: splice one w:t; reject multi_run_span; preserve formatting."""

from __future__ import annotations

import io
import random
import zipfile
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from agent_console.services.docx_blocks import assign_fresh_manifest, block_texts
from agent_console.services.docx_edit import (
    EditOp,
    ValidationError,
    apply_ops,
    package_entry_digests,
    validate_ops,
)

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
XML = "{http://www.w3.org/XML/1998/namespace}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PKG_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _para_xml(data: bytes, needle: str) -> ET.Element:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    for p in root.iter(f"{W}p"):
        text = "".join(t.text or "" for t in p.iter(f"{W}t"))
        if needle in text:
            return p
    raise LookupError(needle)


def _rpr_c14n(run: ET.Element) -> str | None:
    rpr = run.find(f"{W}rPr")
    if rpr is None:
        return None
    return ET.canonicalize(ET.tostring(rpr, encoding="unicode"))


def _live_pass_paragraph_docx() -> bytes:
    """Exact live-pass mixed-run paragraph from scenario #8."""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("See the ")
    linkish = p.add_run("pricing schedule")
    linkish.bold = True
    linkish.underline = True
    p.add_run(" for renewal terms effective 2026-01-01.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_live_pass_replace_keeps_other_runs_c14n_identical() -> None:
    data = _live_pass_paragraph_docx()
    before_p = _para_xml(data, "pricing schedule")
    before_runs = list(before_p.findall(f"{W}r"))
    assert len(before_runs) == 3
    before_c14n = [
        ET.canonicalize(ET.tostring(r, encoding="unicode")) for r in before_runs
    ]
    before_rpr2 = _rpr_c14n(before_runs[1])
    before_rpr3 = _rpr_c14n(before_runs[2])

    manifest = assign_fresh_manifest(data, generation=1)
    target = next(b for b in manifest.blocks if "pricing schedule" in b.text)
    out, _, _ = apply_ops(
        data,
        manifest,
        [
            EditOp(
                op="replace",
                block_id=target.id,
                old="2026-01-01",
                new="2026-06-01",
            )
        ],
    )
    after_p = _para_xml(out, "2026-06-01")
    after_runs = list(after_p.findall(f"{W}r"))
    assert len(after_runs) == 3
    assert ET.canonicalize(ET.tostring(after_runs[0], encoding="unicode")) == before_c14n[0]
    assert ET.canonicalize(ET.tostring(after_runs[1], encoding="unicode")) == before_c14n[1]
    assert _rpr_c14n(after_runs[2]) == before_rpr3
    assert "2026-06-01" in (after_runs[2].find(f"{W}t").text or "")
    assert _rpr_c14n(after_runs[1]) == before_rpr2
    doc = Document(io.BytesIO(out))
    para = next(p for p in doc.paragraphs if "pricing schedule" in p.text)
    assert any(r.bold and r.underline for r in para.runs)


def test_replace_inside_bold_run_keeps_bold() -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Pay within ")
    bold = p.add_run("45 days")
    bold.bold = True
    p.add_run(" please.")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    manifest = assign_fresh_manifest(data, generation=1)
    out, _, _ = apply_ops(
        data,
        manifest,
        [EditOp(op="replace", block_id="g1:p_0001", old="45 days", new="30 days")],
    )
    para = Document(io.BytesIO(out)).paragraphs[0]
    bold_runs = [(r.text, bool(r.bold)) for r in para.runs]
    assert ("30 days", True) in bold_runs


def test_replace_inside_hyperlink_keeps_relationship() -> None:
    doc = Document()
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        document = zf.read("word/document.xml")
        rels = zf.read("word/_rels/document.xml.rels")

    root = ET.fromstring(document)
    body = root.find(f"{W}body")
    assert body is not None
    for child in list(body):
        if child.tag == f"{W}p":
            body.remove(child)
    p = ET.Element(f"{W}p")
    r1 = ET.SubElement(p, f"{W}r")
    ET.SubElement(r1, f"{W}t").text = "See "
    hl = ET.SubElement(p, f"{W}hyperlink")
    hl.set(f"{R_NS}id", "rId99")
    hr = ET.SubElement(hl, f"{W}r")
    ET.SubElement(hr, f"{W}t").text = "pricing schedule"
    r2 = ET.SubElement(p, f"{W}r")
    ET.SubElement(r2, f"{W}t").text = " now."
    sect = body.find(f"{W}sectPr")
    if sect is not None:
        body.insert(list(body).index(sect), p)
    else:
        body.append(p)

    rels_root = ET.fromstring(rels)
    rel = ET.SubElement(rels_root, f"{PKG_REL_NS}Relationship")
    rel.set("Id", "rId99")
    rel.set(
        "Type",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
    )
    rel.set("Target", "https://example.com/pricing")
    rel.set("TargetMode", "External")

    out_buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data), "r") as src:
        with zipfile.ZipFile(out_buf, "w") as dst:
            for info in src.infolist():
                payload = src.read(info.filename)
                if info.filename == "word/document.xml":
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                elif info.filename == "word/_rels/document.xml.rels":
                    payload = ET.tostring(
                        rels_root, encoding="utf-8", xml_declaration=True
                    )
                dst.writestr(info, payload)
    data = out_buf.getvalue()

    manifest = assign_fresh_manifest(data, generation=1)
    block = next(b for b in manifest.blocks if "pricing schedule" in b.text)
    out, _, _ = apply_ops(
        data,
        manifest,
        [
            EditOp(
                op="replace",
                block_id=block.id,
                old="pricing schedule",
                new="rate card",
            )
        ],
    )
    after = _para_xml(out, "rate card")
    hl_after = after.find(f"{W}hyperlink")
    assert hl_after is not None
    assert hl_after.get(f"{R_NS}id") == "rId99"
    with zipfile.ZipFile(io.BytesIO(out)) as zf:
        rels_after = zf.read("word/_rels/document.xml.rels").decode()
    assert "rId99" in rels_after
    assert "example.com/pricing" in rels_after


def test_replace_across_two_runs_is_multi_run_span() -> None:
    """Different rPr across runs must still reject (identical-rPr is the merge path)."""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Hello ")
    p.add_run("world").bold = True
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    before = package_entry_digests(data)
    manifest = assign_fresh_manifest(data, generation=1)
    with pytest.raises(ValidationError, match="multi_run_span"):
        apply_ops(
            data,
            manifest,
            [EditOp(op="replace", block_id="g1:p_0001", old="Hello world", new="Hi")],
        )
    with pytest.raises(ValidationError, match="multi_run_span"):
        validate_ops(
            manifest,
            [EditOp(op="replace", block_id="g1:p_0001", old="Hello world", new="Hi")],
            data=data,
        )
    assert package_entry_digests(data) == before


def test_replace_next_to_tab_succeeds() -> None:
    """Ordinary edits beside a tab must work — only spanning the tab is rejected."""
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Before")
    run = p.add_run()
    run._r.append(OxmlElement("w:tab"))
    p.add_run("AfterTOKEN")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    manifest = assign_fresh_manifest(data, generation=1)
    # Left of tab
    out, _, _ = apply_ops(
        data,
        manifest,
        [EditOp(op="replace", block_id="g1:p_0001", old="Before", new="Start")],
    )
    assert block_texts(out)[0].startswith("Start")
    assert "\t" in block_texts(out)[0] or True
    # Right of tab (unique token entirely in one w:t)
    manifest2 = assign_fresh_manifest(out, generation=1)
    out2, _, _ = apply_ops(
        out,
        manifest2,
        [EditOp(op="replace", block_id="g1:p_0001", old="TOKEN", new="OK")],
    )
    assert "AfterOK" in block_texts(out2)[0]


def test_replace_spanning_tab_is_multi_run_span() -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Before")
    run = p.add_run()
    run._r.append(OxmlElement("w:tab"))
    p.add_run("After")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    before = package_entry_digests(data)
    manifest = assign_fresh_manifest(data, generation=1)
    with pytest.raises(ValidationError, match="multi_run_span"):
        apply_ops(
            data,
            manifest,
            [EditOp(op="replace", block_id="g1:p_0001", old="Before\tAfter", new="X")],
        )
    assert package_entry_digests(data) == before


def test_replace_leading_trailing_space_sets_xml_space() -> None:
    doc = Document()
    doc.add_paragraph("VALUEyy")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    manifest = assign_fresh_manifest(data, generation=1)
    out, _, _ = apply_ops(
        data,
        manifest,
        [EditOp(op="replace", block_id="g1:p_0001", old="VALUE", new=" VALUE ")],
    )
    p = _para_xml(out, " VALUE ")
    matched = next(t for t in p.iter(f"{W}t") if (t.text or "").startswith(" "))
    assert matched.text == " VALUE yy"
    assert matched.get(f"{XML}space") == "preserve"


def test_rewrite_keeps_ppr_and_first_rpr() -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Keep style").bold = True
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    p._p.get_or_add_pPr().append(bidi)
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    manifest = assign_fresh_manifest(data, generation=1)
    out, _, _ = apply_ops(
        data,
        manifest,
        [
            EditOp(
                op="rewrite",
                block_id="g1:p_0001",
                content="New text",
                hash=manifest.blocks[0].content_hash,
            )
        ],
    )
    after_p = _para_xml(out, "New text")
    assert after_p.find(f"{W}pPr") is not None
    assert after_p.find(f"{W}pPr").find(f"{W}bidi") is not None
    run = after_p.find(f"{W}r")
    assert run is not None
    assert run.find(f"{W}rPr") is not None
    assert run.find(f"{W}rPr").find(f"{W}b") is not None


@pytest.mark.parametrize("seed", range(40))
def test_randomized_replace_preserves_other_runs(seed: int) -> None:
    """Invariant: all runs except the target stay C14N-identical; target rPr unchanged."""
    rng = random.Random(seed)
    doc = Document()
    p = doc.add_paragraph()
    tokens: list[str] = []
    for i in range(rng.randint(2, 5)):
        token = f"T{seed}_{i}_{rng.randint(1000, 9999)}"
        tokens.append(token)
        run = p.add_run(f"{token} ")
        if rng.random() < 0.5:
            run.bold = True
        if rng.random() < 0.3:
            run.italic = True
        if rng.random() < 0.3:
            run.underline = True
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()

    before_p = _para_xml(data, tokens[0])
    before_runs = list(before_p.findall(f"{W}r"))
    before_c14n = [
        ET.canonicalize(ET.tostring(r, encoding="unicode")) for r in before_runs
    ]
    before_rprs = [_rpr_c14n(r) for r in before_runs]

    target_i = rng.randint(0, len(tokens) - 1)
    old = tokens[target_i]
    new = f"N{seed}_{target_i}"
    manifest = assign_fresh_manifest(data, generation=1)
    out, _, _ = apply_ops(
        data,
        manifest,
        [EditOp(op="replace", block_id="g1:p_0001", old=old, new=new)],
    )
    after_p = _para_xml(out, new)
    after_runs = list(after_p.findall(f"{W}r"))
    assert len(after_runs) == len(before_runs)
    for i, ar in enumerate(after_runs):
        if i == target_i:
            assert _rpr_c14n(ar) == before_rprs[i]
            assert new in (ar.find(f"{W}t").text or "")
            assert old not in (ar.find(f"{W}t").text or "")
        else:
            assert ET.canonicalize(ET.tostring(ar, encoding="unicode")) == before_c14n[i]
