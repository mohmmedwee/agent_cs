"""Identical-rPr cross-run replace merge (slice 6)."""

from __future__ import annotations

import io
import os
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document

from agent_console.services.docx_blocks import assign_fresh_manifest, block_texts
from agent_console.services.docx_edit import EditOp, ValidationError, apply_ops

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _iter_corpus() -> list[Path]:
    raw = os.environ.get("DOCX_PARITY_DIR", "")
    roots = [Path(p) for p in raw.split(os.pathsep) if p.strip()]
    out: list[Path] = []
    for root in roots:
        if not root.is_dir():
            continue
        for dirpath, dirnames, filenames in os.walk(root):
            rel = Path(dirpath).relative_to(root)
            depth = 0 if str(rel) == "." else len(rel.parts)
            if depth >= 2:
                dirnames.clear()
            for name in filenames:
                if name.endswith(".docx") and not name.startswith("~$"):
                    out.append(Path(dirpath) / name)
    return sorted(out)[:30]


def _para(data: bytes, needle: str) -> ET.Element:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    for p in root.iter(f"{W}p"):
        text = "".join(t.text or "" for t in p.iter(f"{W}t"))
        if needle in text:
            return p
    raise LookupError(needle)


def _docx_from_paragraph_xml(p_xml: str) -> bytes:
    """Wrap a paragraph fragment in a minimal package via python-docx then swap XML."""
    doc = Document()
    doc.add_paragraph("placeholder")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    body = root.find(f"{W}body")
    assert body is not None
    for child in list(body):
        if child.tag == f"{W}p":
            body.remove(child)
    p = ET.fromstring(p_xml)
    sect = body.find(f"{W}sectPr")
    if sect is not None:
        body.insert(list(body).index(sect), p)
    else:
        body.append(p)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data), "r") as src:
        with zipfile.ZipFile(out, "w") as dst:
            for info in src.infolist():
                payload = src.read(info.filename)
                if info.filename == "word/document.xml":
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                dst.writestr(info, payload)
    return out.getvalue()


def test_rsid_only_split_runs_merge() -> None:
    """Same formatting except w:rsid* must merge."""
    p = ET.Element(f"{W}p")
    r1 = ET.SubElement(p, f"{W}r")
    rpr1 = ET.SubElement(r1, f"{W}rPr")
    ET.SubElement(rpr1, f"{W}b")
    ET.SubElement(r1, f"{W}t").text = "Hello "
    r2 = ET.SubElement(p, f"{W}r")
    r2.set(f"{W}rsidR", "00AB12CD")
    rpr2 = ET.SubElement(r2, f"{W}rPr")
    ET.SubElement(rpr2, f"{W}b")
    rpr2.set(f"{W}rsidRPr", "00AB12CD")
    ET.SubElement(r2, f"{W}t").text = "world"
    data = _docx_from_paragraph_xml(ET.tostring(p, encoding="unicode"))
    before = _para(data, "Hello")
    before_outside = [
        ET.canonicalize(ET.tostring(c, encoding="unicode"))
        for c in before
        if c.tag != f"{W}r"
    ]
    manifest = assign_fresh_manifest(data, generation=1)
    out, _, _ = apply_ops(
        data,
        manifest,
        [EditOp(op="replace", block_id="g1:p_0001", old="Hello world", new="Hi")],
    )
    assert block_texts(out)[0] == "Hi"
    after = _para(out, "Hi")
    runs = list(after.findall(f"{W}r"))
    assert len(runs) == 1
    assert runs[0].find(f"{W}rPr").find(f"{W}b") is not None
    after_outside = [
        ET.canonicalize(ET.tostring(c, encoding="unicode"))
        for c in after
        if c.tag != f"{W}r"
    ]
    assert after_outside == before_outside


def test_proof_err_between_runs_merges_and_drops_markers() -> None:
    p = ET.Element(f"{W}p")
    r1 = ET.SubElement(p, f"{W}r")
    ET.SubElement(r1, f"{W}t").text = "Hel"
    err = ET.SubElement(p, f"{W}proofErr")
    err.set(f"{W}type", "spellStart")
    r2 = ET.SubElement(p, f"{W}r")
    ET.SubElement(r2, f"{W}t").text = "lo"
    err2 = ET.SubElement(p, f"{W}proofErr")
    err2.set(f"{W}type", "spellEnd")
    r3 = ET.SubElement(p, f"{W}r")
    ET.SubElement(r3, f"{W}t").text = " there"
    data = _docx_from_paragraph_xml(ET.tostring(p, encoding="unicode"))
    manifest = assign_fresh_manifest(data, generation=1)
    out, _, _ = apply_ops(
        data,
        manifest,
        [EditOp(op="replace", block_id="g1:p_0001", old="Hello", new="Hi")],
    )
    assert block_texts(out)[0] == "Hi there"
    after = _para(out, "Hi")
    # Markers that sat between the merged runs are gone; trailing markers may remain.
    types = [e.get(f"{W}type") for e in after.findall(f"{W}proofErr")]
    assert "spellStart" not in types
    assert len(after.findall(f"{W}r")) == 2  # merged Hello + untouched " there"


def test_different_lang_rejected() -> None:
    p = ET.Element(f"{W}p")
    r1 = ET.SubElement(p, f"{W}r")
    rpr1 = ET.SubElement(r1, f"{W}rPr")
    lang1 = ET.SubElement(rpr1, f"{W}lang")
    lang1.set(f"{W}val", "en-US")
    ET.SubElement(r1, f"{W}t").text = "Hello "
    r2 = ET.SubElement(p, f"{W}r")
    rpr2 = ET.SubElement(r2, f"{W}rPr")
    lang2 = ET.SubElement(rpr2, f"{W}lang")
    lang2.set(f"{W}val", "ar-SA")
    ET.SubElement(r2, f"{W}t").text = "world"
    data = _docx_from_paragraph_xml(ET.tostring(p, encoding="unicode"))
    manifest = assign_fresh_manifest(data, generation=1)
    with pytest.raises(ValidationError, match="multi_run_span"):
        apply_ops(
            data,
            manifest,
            [EditOp(op="replace", block_id="g1:p_0001", old="Hello world", new="Hi")],
        )


def test_bookmark_in_span_rejected() -> None:
    p = ET.Element(f"{W}p")
    r1 = ET.SubElement(p, f"{W}r")
    ET.SubElement(r1, f"{W}t").text = "Hello "
    bm = ET.SubElement(p, f"{W}bookmarkStart")
    bm.set(f"{W}id", "0")
    bm.set(f"{W}name", "_GoBack")
    r2 = ET.SubElement(p, f"{W}r")
    ET.SubElement(r2, f"{W}t").text = "world"
    bm_end = ET.SubElement(p, f"{W}bookmarkEnd")
    bm_end.set(f"{W}id", "0")
    data = _docx_from_paragraph_xml(ET.tostring(p, encoding="unicode"))
    manifest = assign_fresh_manifest(data, generation=1)
    with pytest.raises(ValidationError, match="bookmark"):
        apply_ops(
            data,
            manifest,
            [EditOp(op="replace", block_id="g1:p_0001", old="Hello world", new="Hi")],
        )


def test_merge_leaves_outside_runs_c14n_identical() -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("KEEP ")
    p.add_run("merge")
    p.add_run("me")
    p.add_run(" TAIL")
    # Make middle two identical plain formatting (default)
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    before = _para(data, "KEEP")
    before_runs = list(before.findall(f"{W}r"))
    keep_c14n = ET.canonicalize(ET.tostring(before_runs[0], encoding="unicode"))
    tail_c14n = ET.canonicalize(ET.tostring(before_runs[-1], encoding="unicode"))
    manifest = assign_fresh_manifest(data, generation=1)
    out, _, _ = apply_ops(
        data,
        manifest,
        [EditOp(op="replace", block_id="g1:p_0001", old="mergeme", new="X")],
    )
    after = _para(out, "KEEP")
    after_runs = list(after.findall(f"{W}r"))
    assert ET.canonicalize(ET.tostring(after_runs[0], encoding="unicode")) == keep_c14n
    assert ET.canonicalize(ET.tostring(after_runs[-1], encoding="unicode")) == tail_c14n
    assert "KEEP X TAIL" == block_texts(out)[0]


def test_different_bold_still_multi_run_span() -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Hello ")
    p.add_run("world").bold = True
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    manifest = assign_fresh_manifest(data, generation=1)
    with pytest.raises(ValidationError, match="multi_run_span"):
        apply_ops(
            data,
            manifest,
            [EditOp(op="replace", block_id="g1:p_0001", old="Hello world", new="Hi")],
        )


def test_corpus_cross_run_merge_rate(capsys: pytest.CaptureFixture[str]) -> None:
    """Measure identical-rPr merge vs multi_run_span on DOCX_PARITY_DIR corpus.

    Tries up to three cross-run spans per file (first multi-run paragraphs).
    Prints merge / multi_run_span rates — useful for Phase 2 urgency.
    """
    paths = _iter_corpus()
    if not paths:
        pytest.skip("DOCX_PARITY_DIR not set — no external corpus")
    merged = 0
    rejected = 0
    other = 0
    attempts = 0
    for path in paths:
        try:
            data = path.read_bytes()
            manifest = assign_fresh_manifest(data, generation=1)
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                if "word/document.xml" not in zf.namelist():
                    continue
                root = ET.fromstring(zf.read("word/document.xml"))
            paragraphs = list(root.iter(f"{W}p"))
            file_attempts = 0
            for block in manifest.blocks:
                if file_attempts >= 3:
                    break
                if block.index >= len(paragraphs):
                    continue
                p = paragraphs[block.index]
                runs = [c for c in p if c.tag == f"{W}r"]
                if len(runs) < 2:
                    continue
                run_texts = [
                    "".join((t.text or "") for t in r.iter(f"{W}t")) for r in runs
                ]
                text = block.text
                if len(text) < 4:
                    continue
                t0 = run_texts[0]
                if t0 and len(t0) < len(text):
                    old = text[: len(t0) + min(4, len(text) - len(t0))]
                else:
                    old = text
                if len(old) < 2 or text.count(old) != 1:
                    continue
                attempts += 1
                file_attempts += 1
                try:
                    apply_ops(
                        data,
                        manifest,
                        [
                            EditOp(
                                op="replace",
                                block_id=block.id,
                                old=old,
                                new="X",
                            )
                        ],
                    )
                    merged += 1
                except ValidationError as exc:
                    if "multi_run_span" in str(exc):
                        rejected += 1
                    else:
                        other += 1
                except Exception:  # noqa: BLE001
                    other += 1
        except Exception:  # noqa: BLE001
            other += 1
    total = attempts
    print(
        f"\ncorpus cross-run: attempts={attempts} merged={merged} "
        f"multi_run_span={rejected} other={other}"
    )
    if attempts:
        print(
            f"merge_rate={merged / attempts:.1%} "
            f"reject_rate={rejected / attempts:.1%}"
        )
    assert total > 0
