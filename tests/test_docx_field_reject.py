"""Reject edits to paragraphs with fields or tracked changes."""

from __future__ import annotations

import io
import zipfile
from xml.etree import ElementTree as ET

import pytest
from docx import Document

from agent_console.services.docx_blocks import assign_fresh_manifest
from agent_console.services.docx_edit import EditOp, ValidationError, apply_ops, validate_ops

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx_with_injected(tag: str, attrs: dict[str, str] | None = None) -> bytes:
    """Minimal docx whose first paragraph contains an empty ``tag`` element."""
    doc = Document()
    doc.add_paragraph("Editable field text here")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    body = root.find(f"{W}body")
    assert body is not None
    paragraph = next(p for p in body.iter(f"{W}p") if "".join(
        t.text or "" for t in p.iter(f"{W}t")
    ).strip())
    el = ET.SubElement(paragraph, tag)
    for key, value in (attrs or {}).items():
        el.set(key, value)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data), "r") as src:
        with zipfile.ZipFile(out, "w") as dst:
            for info in src.infolist():
                payload = src.read(info.filename)
                if info.filename == "word/document.xml":
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                dst.writestr(info, payload)
    return out.getvalue()


@pytest.mark.parametrize(
    "tag,local",
    [
        (f"{W}fldChar", "fldChar"),
        (f"{W}fldSimple", "fldSimple"),
        (f"{W}ins", "ins"),
        (f"{W}del", "del"),
    ],
)
def test_replace_rejects_paragraph_with_field_or_revision(tag: str, local: str) -> None:
    attrs = {f"{W}fldCharType": "begin"} if local == "fldChar" else None
    data = _docx_with_injected(tag, attrs)
    manifest = assign_fresh_manifest(data, generation=1)
    block = manifest.blocks[0]
    with pytest.raises(ValidationError, match=local):
        validate_ops(
            manifest,
            [
                EditOp(
                    op="replace",
                    block_id=block.id,
                    old="Editable",
                    new="Changed",
                )
            ],
            data=data,
        )
    with pytest.raises(ValidationError, match=local):
        apply_ops(
            data,
            manifest,
            [
                EditOp(
                    op="replace",
                    block_id=block.id,
                    old="Editable",
                    new="Changed",
                )
            ],
        )


def test_plain_paragraph_still_editable() -> None:
    doc = Document()
    doc.add_paragraph("Plain text only")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    manifest = assign_fresh_manifest(data, generation=1)
    out, _, _ = apply_ops(
        data,
        manifest,
        [EditOp(op="replace", block_id="g1:p_0001", old="Plain", new="Clean")],
    )
    assert "Clean text only" in "".join(
        __import__("agent_console.services.docx_blocks", fromlist=["block_texts"]).block_texts(
            out
        )
    )
