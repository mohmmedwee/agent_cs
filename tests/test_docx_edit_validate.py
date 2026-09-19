"""Phase 1 slice 3: edit_docx validator (dry-run) + plaintext diff."""

from __future__ import annotations

import io

import pytest
from docx import Document

from agent_console.services.docx_blocks import assign_fresh_manifest
from agent_console.services.docx_edit import (
    EditOp,
    ValidationError,
    validate_ops,
)


def _minimal_docx(*paragraphs: str) -> bytes:
    doc = Document()
    if doc.paragraphs:
        doc.paragraphs[0].text = paragraphs[0] if paragraphs else ""
        rest = paragraphs[1:]
    else:
        rest = paragraphs
    for text in rest:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_replace_exact_once_produces_diff() -> None:
    data = _minimal_docx("Due in 45 days.", "Other.")
    manifest = assign_fresh_manifest(data, generation=1)
    result = validate_ops(
        manifest,
        [EditOp(op="replace", block_id="g1:p_0001", old="45 days", new="30 days")],
    )
    assert result.ok
    assert len(result.diffs) == 1
    d = result.diffs[0]
    assert d.block_id == "g1:p_0001"
    assert d.before == "Due in 45 days."
    assert d.after == "Due in 30 days."


def test_replace_zero_matches_errors() -> None:
    data = _minimal_docx("Due in 45 days.")
    manifest = assign_fresh_manifest(data, generation=1)
    with pytest.raises(ValidationError, match="0 matches"):
        validate_ops(
            manifest,
            [EditOp(op="replace", block_id="g1:p_0001", old="99 days", new="30")],
        )


def test_replace_two_matches_without_count_errors() -> None:
    data = _minimal_docx("aa aa")
    manifest = assign_fresh_manifest(data, generation=1)
    with pytest.raises(ValidationError, match="2 matches"):
        validate_ops(
            manifest,
            [EditOp(op="replace", block_id="g1:p_0001", old="aa", new="bb")],
        )


def test_wrong_generation_asks_reread() -> None:
    data = _minimal_docx("Hello")
    manifest = assign_fresh_manifest(data, generation=1)
    with pytest.raises(ValidationError, match="re-read"):
        validate_ops(
            manifest,
            [EditOp(op="replace", block_id="g2:p_0001", old="Hello", new="Hi")],
        )


def test_rewrite_requires_matching_hash() -> None:
    data = _minimal_docx("Hello")
    manifest = assign_fresh_manifest(data, generation=1)
    b = manifest.blocks[0]
    with pytest.raises(ValidationError, match="hash"):
        validate_ops(
            manifest,
            [
                EditOp(
                    op="rewrite",
                    block_id=b.id,
                    content="Hi",
                    hash="deadbeef",
                )
            ],
        )
    result = validate_ops(
        manifest,
        [
            EditOp(
                op="rewrite",
                block_id=b.id,
                content="Hi",
                hash=b.content_hash,
            )
        ],
    )
    assert result.diffs[0].after == "Hi"


def test_delete_requires_matching_hash() -> None:
    data = _minimal_docx("Keep", "Drop")
    manifest = assign_fresh_manifest(data, generation=1)
    drop = manifest.blocks[1]
    result = validate_ops(
        manifest,
        [EditOp(op="delete", block_id=drop.id, hash=drop.content_hash)],
    )
    assert result.diffs[0].before == "Drop"
    assert result.diffs[0].after == ""


def test_insert_relative_to_wrong_generation_rejected() -> None:
    data = _minimal_docx("Anchor")
    manifest = assign_fresh_manifest(data, generation=1)
    with pytest.raises(ValidationError, match="re-read"):
        validate_ops(
            manifest,
            [
                EditOp(
                    op="insert",
                    relative_to="g9:p_0001",
                    position="after",
                    content="New",
                )
            ],
        )


def test_insert_after_produces_diff_with_real_id() -> None:
    data = _minimal_docx("Anchor")
    manifest = assign_fresh_manifest(data, generation=1)
    result = validate_ops(
        manifest,
        [
            EditOp(
                op="insert",
                relative_to="g1:p_0001",
                position="after",
                content="Inserted",
            )
        ],
    )
    assert result.ok
    assert result.diffs[0].block_id == "g1:p_0002"
    assert result.temp_id_map == {"new_1": "g1:p_0002"}
    assert result.diffs[0].after == "Inserted"
    assert result.diffs[0].before == ""
