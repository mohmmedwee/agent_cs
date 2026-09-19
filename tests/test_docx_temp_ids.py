"""Temp-ID insert chains in validator and executor."""

from __future__ import annotations

import io

import pytest
from docx import Document

from agent_console.services.docx_blocks import assign_fresh_manifest, block_texts
from agent_console.services.docx_edit import EditOp, ValidationError, apply_ops, validate_ops


def _minimal_docx(*paragraphs: str) -> bytes:
    doc = Document()
    if doc.paragraphs:
        doc.paragraphs[0].text = paragraphs[0] if paragraphs else ""
        rest = paragraphs[1:]
    else:
        rest = paragraphs
    for text in rest:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_live_pass_temp_id_insert_chain() -> None:
    """Exact live payload from scenario #4."""
    data = _minimal_docx("Service Agreement", "Body")
    manifest = assign_fresh_manifest(data, generation=1)
    assert manifest.next_id == 3
    ops = [
        EditOp(
            op="insert",
            relative_to="g1:p_0001",
            position="before",
            content="NEW HEADING",
            new_id="new_1",
        ),
        EditOp(
            op="insert",
            relative_to="new_1",
            position="after",
            content="P1",
            new_id="new_2",
        ),
        EditOp(
            op="insert",
            relative_to="new_2",
            position="after",
            content="P2",
            new_id="new_3",
        ),
    ]
    result = validate_ops(manifest, ops, data=data)
    assert result.temp_id_map == {
        "new_1": "g1:p_0003",
        "new_2": "g1:p_0004",
        "new_3": "g1:p_0005",
    }
    out, new_manifest, diffs = apply_ops(data, manifest, ops)
    assert [d.block_id for d in diffs] == ["g1:p_0003", "g1:p_0004", "g1:p_0005"]
    assert block_texts(out)[:4] == [
        "NEW HEADING",
        "P1",
        "P2",
        "Service Agreement",
    ]
    assert [b.id for b in new_manifest.blocks[:3]] == [
        "g1:p_0003",
        "g1:p_0004",
        "g1:p_0005",
    ]
    assert new_manifest.next_id == 6
    by_id = {b.id: b.text for b in new_manifest.blocks}
    assert by_id["g1:p_0003"] == "NEW HEADING"
    assert by_id["g1:p_0004"] == "P1"
    assert by_id["g1:p_0005"] == "P2"


def test_forward_temp_reference_rejected() -> None:
    data = _minimal_docx("Anchor")
    manifest = assign_fresh_manifest(data, generation=1)
    with pytest.raises(ValidationError, match="forward|unknown temp"):
        validate_ops(
            manifest,
            [
                EditOp(
                    op="insert",
                    relative_to="new_2",
                    position="after",
                    content="Too early",
                    new_id="new_1",
                )
            ],
            data=data,
        )


def test_duplicate_temp_id_rejected() -> None:
    data = _minimal_docx("Anchor")
    manifest = assign_fresh_manifest(data, generation=1)
    with pytest.raises(ValidationError, match="duplicate temp"):
        validate_ops(
            manifest,
            [
                EditOp(
                    op="insert",
                    relative_to="g1:p_0001",
                    position="after",
                    content="A",
                    new_id="new_1",
                ),
                EditOp(
                    op="insert",
                    relative_to="new_1",
                    position="after",
                    content="B",
                    new_id="new_1",
                ),
            ],
            data=data,
        )


def test_temp_id_as_delete_target_rejected() -> None:
    data = _minimal_docx("Anchor")
    manifest = assign_fresh_manifest(data, generation=1)
    with pytest.raises(ValidationError, match="cannot be the target"):
        validate_ops(
            manifest,
            [
                EditOp(
                    op="insert",
                    relative_to="g1:p_0001",
                    position="after",
                    content="A",
                    new_id="new_1",
                ),
                EditOp(op="delete", block_id="new_1", hash="deadbeef"),
            ],
            data=data,
        )


def test_generation_prefixed_temp_id_rejected() -> None:
    data = _minimal_docx("Anchor")
    manifest = assign_fresh_manifest(data, generation=1)
    with pytest.raises(ValidationError, match="don't take a generation"):
        validate_ops(
            manifest,
            [
                EditOp(
                    op="insert",
                    relative_to="g1:p_0001",
                    position="after",
                    content="A",
                    new_id="new_1",
                ),
                EditOp(
                    op="insert",
                    relative_to="g1:new_1",
                    position="after",
                    content="B",
                    new_id="new_2",
                ),
            ],
            data=data,
        )
