"""Carry-forward manifests and generation bumps after convert / run_python."""

from __future__ import annotations

import io
from unittest.mock import MagicMock

import pytest
from docx import Document

from agent_console.config import get_settings
from agent_console.db.models import User
from agent_console.repositories.documents import DOCX_MEDIA_TYPE
from agent_console.repositories.files import FileRepository
from agent_console.services.docx_blocks import assign_fresh_manifest, block_texts
from agent_console.services.docx_edit import EditOp, ValidationError, apply_ops, validate_ops
from agent_console.services.tools import convert_upload as convert_mod
from agent_console.services.tools import run_python as run_python_mod
from agent_console.services.tools.context import ToolContext
from agent_console.services.tools.registry import ToolRegistry


def _minimal_docx(*paragraphs: str) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _tool_registry(files: FileRepository, user: User) -> ToolRegistry:
    registry = ToolRegistry()
    ctx = ToolContext(
        settings=get_settings(),
        files=files,
        memories=MagicMock(),
        skills=MagicMock(),
        search=MagicMock(),
        pages=MagicMock(),
        cache=MagicMock(),
        user_id=user.id,
        upstream=MagicMock(),
    )
    convert_mod.register(registry, ctx)
    run_python_mod.register(registry, ctx)
    return registry


def test_insert_above_then_edit_keeps_old_id_on_original_paragraph() -> None:
    """Carry-forward: insert shifts order but g1:p_0001 still names the old text."""
    data = _minimal_docx("Alpha", "Beta", "Gamma")
    manifest = assign_fresh_manifest(data, generation=1)
    alpha_id = "g1:p_0001"
    assert manifest.blocks[0].text == "Alpha"

    out, carried, _ = apply_ops(
        data,
        manifest,
        [
            EditOp(
                op="insert",
                relative_to=alpha_id,
                position="before",
                content="NEW",
                new_id="new_1",
            )
        ],
    )
    assert block_texts(out)[:2] == ["NEW", "Alpha"]
    by_id = {b.id: b for b in carried.blocks}
    assert by_id[alpha_id].text == "Alpha"
    assert by_id[alpha_id].index == 1

    out2, carried2, diffs = apply_ops(
        out,
        carried,
        [EditOp(op="replace", block_id=alpha_id, old="Alpha", new="ALPHA")],
    )
    assert diffs[0].after == "ALPHA"
    assert block_texts(out2)[:2] == ["NEW", "ALPHA"]
    assert {b.id: b.text for b in carried2.blocks}[alpha_id] == "ALPHA"


def test_fresh_reassign_after_insert_would_point_at_wrong_paragraph() -> None:
    """Documents why persisted manifests matter: reassign maps p_0001 to NEW."""
    data = _minimal_docx("Alpha", "Beta")
    manifest = assign_fresh_manifest(data, generation=1)
    out, _, _ = apply_ops(
        data,
        manifest,
        [
            EditOp(
                op="insert",
                relative_to="g1:p_0001",
                position="before",
                content="NEW",
            )
        ],
    )
    reassigned = assign_fresh_manifest(out, generation=1)
    assert reassigned.blocks[0].text == "NEW"
    assert reassigned.blocks[0].id == "g1:p_0001"


def test_bumped_generation_rejects_old_ids_with_reread() -> None:
    """When generation advances, stale gN ids fail closed."""
    data = _minimal_docx("Hello")
    m2 = assign_fresh_manifest(data, generation=2)
    with pytest.raises(ValidationError, match="re-read"):
        validate_ops(
            m2,
            [EditOp(op="replace", block_id="g1:p_0001", old="Hello", new="Hi")],
            data=data,
        )


@pytest.mark.db
async def test_after_convert_rebuild_bumps_generation(
    files: FileRepository, user: User
) -> None:
    await files.save(user.id, name="src.md", data=b"# Title\n\nHello body\n")
    registry = _tool_registry(files, user)
    first = await registry.invoke(
        "convert_upload_to_docx",
        '{"source":"src.md","name":"out.docx"}',
    )
    assert first.startswith("Converted")
    tip1 = await files.latest_in_chain(user.id, "out.docx")
    assert tip1.block_generation == 1
    second = await registry.invoke(
        "convert_upload_to_docx",
        '{"source":"src.md","name":"out.docx"}',
    )
    assert "Converted" in second
    tip2 = await files.latest_in_chain(user.id, "out.docx")
    assert tip2.block_generation == 2
    data = await files.raw_bytes(user.id, str(tip2.id))
    m2 = assign_fresh_manifest(data, generation=tip2.block_generation)
    with pytest.raises(ValidationError, match="re-read"):
        validate_ops(
            m2,
            [EditOp(op="replace", block_id="g1:p_0001", old="Hello", new="X")],
            data=data,
        )


@pytest.mark.db
async def test_after_run_python_one_docx_input_bumps_generation(
    files: FileRepository, user: User
) -> None:
    tip = await files.save(
        user.id,
        name="in.docx",
        data=_minimal_docx("Script target"),
        content_type=DOCX_MEDIA_TYPE,
        block_generation=1,
    )
    assert tip.block_generation == 1
    registry = _tool_registry(files, user)
    code = (
        "from docx import Document\n"
        "d = Document('in.docx')\n"
        "d.paragraphs[0].text = 'Script rewritten body'\n"
        "d.save('out.docx')\n"
    )
    result = await registry.invoke(
        "run_python",
        f'{{"code":{code!r},"inputs":["in.docx"]}}',
    )
    assert "out.docx" in result
    out = await files.get(user.id, "out.docx")
    assert out.parent_id == tip.id
    assert out.block_generation == 2
    data = await files.raw_bytes(user.id, str(out.id))
    m = assign_fresh_manifest(data, generation=out.block_generation)
    with pytest.raises(ValidationError, match="re-read"):
        validate_ops(
            m,
            [EditOp(op="replace", block_id="g1:p_0001", old="x", new="y")],
            data=data,
        )


@pytest.mark.db
async def test_run_python_two_docx_inputs_creates_new_root(
    files: FileRepository, user: User
) -> None:
    await files.save(
        user.id, name="a.docx", data=_minimal_docx("A"), content_type=DOCX_MEDIA_TYPE
    )
    await files.save(
        user.id, name="b.docx", data=_minimal_docx("B"), content_type=DOCX_MEDIA_TYPE
    )
    registry = _tool_registry(files, user)
    code = (
        "from docx import Document\n"
        "d = Document()\n"
        "d.add_paragraph('merged')\n"
        "d.save('merged.docx')\n"
    )
    await registry.invoke(
        "run_python",
        f'{{"code":{code!r},"inputs":["a.docx","b.docx"]}}',
    )
    merged = await files.get(user.id, "merged.docx")
    assert merged.parent_id is None
    assert merged.block_generation == 1
    assert merged.root_id == merged.id
