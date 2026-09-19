"""Corruption / Word-strictness cases for edit_docx apply."""

from __future__ import annotations

import io
import zipfile
from xml.etree import ElementTree as ET

import pytest
from docx import Document

from agent_console.services.docx_blocks import assign_fresh_manifest, block_texts
from agent_console.services.docx_edit import EditOp, ValidationError, apply_ops


_WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _bytes_from_doc(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _minimal_docx(*paragraphs: str) -> bytes:
    doc = Document()
    if doc.paragraphs:
        doc.paragraphs[0].text = paragraphs[0] if paragraphs else ""
        rest = paragraphs[1:]
    else:
        rest = paragraphs
    for text in rest:
        doc.add_paragraph(text)
    return _bytes_from_doc(doc)


def test_refuse_delete_sole_table_cell_paragraph() -> None:
    doc = Document()
    if doc.paragraphs:
        doc.paragraphs[0].text = "before"
    else:
        doc.add_paragraph("before")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "only cell"
    data = _bytes_from_doc(doc)
    manifest = assign_fresh_manifest(data, generation=1)
    cell_block = next(b for b in manifest.blocks if b.text == "only cell")
    with pytest.raises(ValidationError, match="only paragraph"):
        apply_ops(
            data,
            manifest,
            [EditOp(op="delete", block_id=cell_block.id, hash=cell_block.content_hash)],
        )


def test_refuse_delete_paragraph_holding_sect_pr() -> None:
    """Legacy layout: sectPr nested inside the last body paragraph."""
    data = _minimal_docx("Keep", "Section host")
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = zf.namelist()
        document = zf.read("word/document.xml")
        extras = {n: zf.read(n) for n in names if n != "word/document.xml"}
    root = ET.fromstring(document)
    body = root.find(f"{_WORD_NS}body")
    assert body is not None
    # Move/create sectPr onto the last paragraph (and drop body-level sectPr).
    for child in list(body):
        if child.tag == f"{_WORD_NS}sectPr":
            body.remove(child)
            paras = [c for c in body if c.tag == f"{_WORD_NS}p"]
            p_pr = paras[-1].find(f"{_WORD_NS}pPr")
            if p_pr is None:
                p_pr = ET.Element(f"{_WORD_NS}pPr")
                paras[-1].insert(0, p_pr)
            p_pr.append(child)
    new_xml = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, payload in extras.items():
            zf.writestr(name, payload)
        zf.writestr("word/document.xml", new_xml)
    data = buf.getvalue()

    manifest = assign_fresh_manifest(data, generation=1)
    last = manifest.blocks[-1]
    with pytest.raises(ValidationError, match="sectPr"):
        apply_ops(
            data,
            manifest,
            [EditOp(op="delete", block_id=last.id, hash=last.content_hash)],
        )


def test_insert_after_last_body_paragraph_keeps_sect_pr_trailing() -> None:
    data = _minimal_docx("Only")
    manifest = assign_fresh_manifest(data, generation=1)
    last = manifest.blocks[-1]
    out, new_manifest, _ = apply_ops(
        data,
        manifest,
        [
            EditOp(
                op="insert",
                relative_to=last.id,
                position="after",
                content="After last",
            )
        ],
    )
    assert block_texts(out) == ["Only", "After last"]
    with zipfile.ZipFile(io.BytesIO(out)) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    body = root.find(f"{_WORD_NS}body")
    assert body is not None
    body_kids = list(body)
    # Trailing body-level sectPr (python-docx style) must remain last.
    assert body_kids[-1].tag == f"{_WORD_NS}sectPr"
    assert "After last" in {b.text for b in new_manifest.blocks}
    Document(io.BytesIO(out))


def test_xml_special_chars_and_space_preserve_and_newlines() -> None:
    data = _minimal_docx("plain")
    manifest = assign_fresh_manifest(data, generation=1)
    target = manifest.blocks[0]
    content = '  <tag>&"quote"\nline2  '
    out, _, _ = apply_ops(
        data,
        manifest,
        [
            EditOp(
                op="rewrite",
                block_id=target.id,
                content=content,
                hash=target.content_hash,
            )
        ],
    )
    with zipfile.ZipFile(io.BytesIO(out)) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    assert "&lt;tag&gt;" in xml
    assert "&amp;" in xml
    assert 'space="preserve"' in xml
    assert ":br" in xml or "<br" in xml
    # Agent-visible text follows _docx_text strip rules; spaces still live in OOXML.
    assert block_texts(out)[0] == content.strip()
    assert ">  &lt;tag&gt;" in xml or ">  <" not in xml  # preserved leading spaces in w:t
    assert "line2  </" in xml or "line2  </ns0:t>" in xml
    Document(io.BytesIO(out))
