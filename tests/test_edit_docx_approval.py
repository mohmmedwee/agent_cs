"""edit_docx approval: deny / stale tip revalidation."""

from __future__ import annotations

import io

import pytest
from docx import Document

from agent_console.db.models import User
from agent_console.repositories.files import FileRepository
from agent_console.services.approvals import ApprovalBroker
from agent_console.services.docx_blocks import assign_fresh_manifest
from agent_console.services.tools.edit_docx import (
    apply_edit_docx_payload,
    prepare_edit_docx_approval,
)


pytestmark = pytest.mark.db


def _minimal_docx(*paragraphs: str) -> bytes:
    doc = Document()
    if doc.paragraphs:
        doc.paragraphs[0].text = paragraphs[0] if paragraphs else ""
        rest = paragraphs[1:]
    else:
        rest = paragraphs
    for text in rest:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def test_deny_produces_no_new_version(
    files: FileRepository, user: User
) -> None:
    tip = await files.save(
        user.id, name="deny.docx", data=_minimal_docx("Hello"), content_type="docx"
    )
    payload, card, err = await prepare_edit_docx_approval(
        files,
        user.id,
        '{"name":"deny.docx","operations":[{"op":"replace","block_id":"g1:p_0001","old":"Hello","new":"Hi"}]}',
    )
    assert err is None and payload and card
    broker = ApprovalBroker()
    broker.put_payload("c", "call", payload, ttl_seconds=60)
    # Deny path clears payload without apply.
    taken = broker.take_payload("c", "call")
    assert taken is not None
    # Simulate deny: discard without applying.
    tip2 = await files.latest_in_chain(user.id, "deny.docx")
    assert tip2.id == tip.id
    assert tip2.version == 1


async def test_allow_with_stale_tip_shows_clear_error(
    files: FileRepository, user: User
) -> None:
    data = _minimal_docx("Hello world")
    tip = await files.save(user.id, name="stale.docx", data=data)
    payload, _, err = await prepare_edit_docx_approval(
        files,
        user.id,
        '{"name":"stale.docx","operations":[{"op":"replace","block_id":"g1:p_0001","old":"world","new":"there"}]}',
    )
    assert err is None and payload
    # Tip moves before apply (another edit wins the CAS parent).
    await files.save(
        user.id,
        name="stale.docx",
        data=data,
        parent_id=tip.id,
        derived_from=tip.id,
    )
    result = await apply_edit_docx_payload(files, user.id, payload)
    assert result.startswith("Error:")
    assert "nothing applied" in result.lower() or "changed" in result.lower()
    tip2 = await files.latest_in_chain(user.id, "stale.docx")
    assert tip2.version == 2
    # Body unchanged from the intermediate tip (still original bytes).
    assert await files.raw_bytes(user.id, str(tip2.id)) == data
