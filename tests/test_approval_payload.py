"""ApprovalBroker stores only plain JSON payloads (Redis-ready, single-use)."""

from __future__ import annotations

import asyncio
import io
import time

import pytest

from agent_console.services.approvals import ApprovalBroker, _as_json_dict
from agent_console.services.docx_blocks import assign_fresh_manifest
from agent_console.services.docx_edit import (
    ValidationError,
    apply_approved_payload,
)
from docx import Document


def _minimal_docx(*paragraphs: str) -> bytes:
    doc = Document()
    if doc.paragraphs:
        doc.paragraphs[0].text = paragraphs[0] if paragraphs else ""
        rest = paragraphs[1:]
    else:
        rest = paragraphs
    for text in rest:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_payload_roundtrip_keeps_ops_diff_and_ids() -> None:
    payload = {
        "source_file_id": "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee",
        "generation": 1,
        "hashes": {"g1:p_0001": "a1b2c3d4"},
        "operations": [
            {
                "op": "replace",
                "block_id": "g1:p_0001",
                "old": "45 days",
                "new": "30 days",
            }
        ],
        "diff_text": "### g1:p_0001\n- Due in 45 days.\n+ Due in 30 days.",
    }
    clean = _as_json_dict(payload)
    assert clean == payload


def test_payload_rejects_non_json_objects() -> None:
    with pytest.raises(TypeError, match="JSON-serializable"):
        _as_json_dict({"doc": object()})


async def test_wait_keeps_payload_after_allow_for_apply() -> None:
    broker = ApprovalBroker()
    payload = {
        "source_file_id": "file-1",
        "generation": 2,
        "hashes": {"g2:p_0001": "deadbeef"},
        "operations": [{"op": "delete", "block_id": "g2:p_0001", "hash": "deadbeef"}],
        "diff_text": "### g2:p_0001\n- gone\n+ ",
    }

    async def allow_soon() -> None:
        await asyncio.sleep(0.05)
        assert broker.get_payload("c1", "call1") == payload
        assert broker.resolve("c1", "call1", True)

    task = asyncio.create_task(allow_soon())
    allowed = await broker.wait("c1", "call1", timeout=2.0, payload=payload)
    await task
    assert allowed is True
    assert broker.take_payload("c1", "call1") == payload
    assert broker.get_payload("c1", "call1") is None


async def test_wait_clears_payload_on_deny() -> None:
    broker = ApprovalBroker()
    payload = {
        "source_file_id": "file-1",
        "generation": 1,
        "hashes": {},
        "operations": [{"op": "rewrite", "block_id": "g1:p_0001", "content": "x", "hash": "a"}],
        "diff_text": "",
    }

    async def deny_soon() -> None:
        await asyncio.sleep(0.05)
        broker.resolve("c1", "call2", False)

    task = asyncio.create_task(deny_soon())
    allowed = await broker.wait("c1", "call2", timeout=2.0, payload=payload)
    await task
    assert allowed is False
    assert broker.get_payload("c1", "call2") is None


def test_take_payload_is_single_use() -> None:
    broker = ApprovalBroker()
    payload = {
        "source_file_id": "f",
        "generation": 1,
        "hashes": {},
        "operations": [
            {
                "op": "insert",
                "relative_to": "g1:p_0001",
                "position": "after",
                "content": "X",
            }
        ],
        "diff_text": "",
    }
    broker.put_payload("c", "call", payload, ttl_seconds=60)
    assert broker.take_payload("c", "call") == payload
    assert broker.take_payload("c", "call") is None


def test_payload_expires_by_ttl() -> None:
    broker = ApprovalBroker()
    broker.put_payload(
        "c",
        "call",
        {
            "source_file_id": "f",
            "generation": 1,
            "hashes": {},
            "operations": [{"op": "insert", "relative_to": "g1:p_0001", "content": "X"}],
            "diff_text": "",
        },
        ttl_seconds=0.05,
    )
    time.sleep(0.08)
    assert broker.get_payload("c", "call") is None
    assert broker.take_payload("c", "call") is None


def test_apply_same_payload_twice_rejects_second_insert() -> None:
    """Insert replay would pass revalidation; single-use payload must block it."""
    data = _minimal_docx("Anchor")
    manifest = assign_fresh_manifest(data, generation=1)
    broker = ApprovalBroker()
    payload = {
        "source_file_id": "file-1",
        "generation": 1,
        "hashes": {manifest.blocks[0].id: manifest.blocks[0].content_hash},
        "operations": [
            {
                "op": "insert",
                "relative_to": "g1:p_0001",
                "position": "after",
                "content": "Inserted once",
            }
        ],
        "diff_text": "### new_1\n- \n+ Inserted once",
    }
    broker.put_payload("conv", "call-insert", payload, ttl_seconds=120)

    out, new_manifest, _ = apply_approved_payload(
        broker=broker,
        conversation_id="conv",
        call_id="call-insert",
        data=data,
        manifest=manifest,
    )
    assert any(b.text == "Inserted once" for b in new_manifest.blocks)

    with pytest.raises(ValidationError, match="already applied"):
        apply_approved_payload(
            broker=broker,
            conversation_id="conv",
            call_id="call-insert",
            data=out,
            manifest=new_manifest,
        )
