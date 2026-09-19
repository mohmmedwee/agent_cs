"""Phase 1 slice 2: annotated DOCX read/search lines."""

from __future__ import annotations

import io

from docx import Document

from agent_console.services.docx_blocks import (
    assign_fresh_manifest,
    format_annotated,
    format_block_line,
    search_blocks,
)


def _minimal_docx(*paragraphs: str) -> bytes:
    doc = Document()
    if doc.paragraphs:
        doc.paragraphs[0].text = paragraphs[0] if paragraphs else ""
        rest = paragraphs[1:]
    else:
        rest = paragraphs
    for text in rest:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_format_block_line_uses_generation_id_and_hash() -> None:
    data = _minimal_docx("Payment due in 45 days.")
    manifest = assign_fresh_manifest(data, generation=1)
    line = format_block_line(manifest.blocks[0])
    b = manifest.blocks[0]
    assert line == f"[{b.id} h={b.content_hash}] {b.text}"
    assert line.startswith("[g1:p_0001 h=")


def test_format_annotated_one_line_per_block() -> None:
    data = _minimal_docx("Alpha", "Beta")
    manifest = assign_fresh_manifest(data, generation=2)
    text = format_annotated(manifest)
    lines = text.splitlines()
    assert len(lines) == 2
    assert lines[0].startswith("[g2:p_0001 h=")
    assert lines[0].endswith("Alpha")
    assert lines[1].startswith("[g2:p_0002 h=")
    assert lines[1].endswith("Beta")


def test_search_blocks_returns_matching_ids_and_local_offsets() -> None:
    data = _minimal_docx("Hello world", "No match here", "world peace")
    manifest = assign_fresh_manifest(data, generation=1)
    hits = search_blocks(manifest, "world")
    assert [h.block_id for h in hits] == ["g1:p_0001", "g1:p_0003"]
    assert hits[0].start == 6
    assert hits[0].end == 11
    assert "world" in hits[0].excerpt.lower()
