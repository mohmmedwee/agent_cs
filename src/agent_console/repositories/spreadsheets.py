"""Building a real .xlsx from Markdown (or CSV) the model writes.

Same idea as `documents.build_docx`: the user asked for a spreadsheet, not a
text file wearing an `.xlsx` name. Tables in the Markdown become sheets;
plain CSV content becomes one sheet; leftover prose lands on a Notes sheet.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

__all__ = ["XLSX_MEDIA_TYPE", "build_xlsx"]

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_TABLE_SPLIT = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")


def _sheet_title(raw: str | None, index: int) -> str:
    name = (raw or f"Sheet{index}").strip() or f"Sheet{index}"
    # Excel sheet names: max 31 chars, no []:*?/\
    cleaned = re.sub(r'[\[\]:*?/\\]', "-", name)[:31]
    return cleaned or f"Sheet{index}"


def _cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def _is_rule(line: str) -> bool:
    return bool(_TABLE_SPLIT.match(line.strip()))


def _parse_markdown_tables(markdown: str) -> list[tuple[str | None, list[str], list[list[str]]]]:
    """Return (heading_or_none, header_row, data_rows) for each GFM pipe table."""
    lines = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    tables: list[tuple[str | None, list[str], list[list[str]]]] = []
    last_heading: str | None = None
    index = 0
    while index < len(lines):
        line = lines[index]
        heading = re.match(r"^(#{1,6})\s+(.+)$", line.strip())
        if heading:
            last_heading = heading.group(2).strip()
            index += 1
            continue

        if (
            "|" in line
            and index + 1 < len(lines)
            and _is_rule(lines[index + 1])
        ):
            header = _cells(line)
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                if _is_rule(lines[index]):
                    index += 1
                    continue
                rows.append(_cells(lines[index]))
                index += 1
            # Normalize row widths to header.
            width = len(header)
            normalized = [
                (row + [""] * width)[:width] for row in rows if any(cell.strip() for cell in row)
            ]
            tables.append((last_heading, header, normalized))
            last_heading = None
            continue

        index += 1
    return tables


def _looks_like_csv(text: str) -> bool:
    sample = text.strip()
    if not sample or "|" in sample.split("\n", 1)[0]:
        return False
    try:
        dialect = csv.Sniffer().sniff(sample[:4000], delimiters=",;\t")
        rows = list(csv.reader(io.StringIO(sample), dialect))
    except csv.Error:
        return False
    return len(rows) >= 2 and max(len(row) for row in rows) >= 2


def _parse_csv(text: str) -> tuple[list[str], list[list[str]]]:
    sample = text.strip()
    try:
        dialect = csv.Sniffer().sniff(sample[:4000], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(sample), dialect))
    if not rows:
        return [], []
    header = [cell.strip() or f"Column{i + 1}" for i, cell in enumerate(rows[0])]
    body = [row for row in rows[1:] if any(cell.strip() for cell in row)]
    width = len(header)
    body = [(row + [""] * width)[:width] for row in body]
    return header, body


def _prose_lines(markdown: str, skip_tables: bool = True) -> list[str]:
    """Non-table lines useful as Notes (skip front matter and table blocks)."""
    text = markdown.replace("\r\n", "\n").replace("\r", "\n")
    if text.startswith("---"):
        parts = text.split("\n")
        if len(parts) > 1:
            for end, line in enumerate(parts[1:], start=1):
                if line.strip() == "---":
                    text = "\n".join(parts[end + 1 :])
                    break

    if not skip_tables:
        return [line for line in text.split("\n") if line.strip()]

    lines = text.split("\n")
    kept: list[str] = []
    index = 0
    while index < len(lines):
        line = lines[index]
        if (
            "|" in line
            and index + 1 < len(lines)
            and _is_rule(lines[index + 1])
        ):
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                index += 1
            continue
        if line.strip():
            kept.append(line.rstrip())
        index += 1
    return kept


def _cell_value(text: str) -> tuple[str, bool]:
    """Strip common inline Markdown; return (plain_text, wants_bold)."""
    value = (text or "").strip()
    bold = False

    def _unwrap(pattern: str, mark_bold: bool = False) -> None:
        nonlocal value, bold
        match = re.fullmatch(pattern, value, flags=re.DOTALL)
        if match:
            value = match.group(1)
            if mark_bold:
                bold = True

    _unwrap(r"\*\*(.+)\*\*", mark_bold=True)
    _unwrap(r"__(.+)__", mark_bold=True)
    _unwrap(r"\*(.+)\*", mark_bold=True)
    _unwrap(r"_(.+)_", mark_bold=True)
    _unwrap(r"`(.+)`")
    return value, bold


def _write_sheet(ws, header: list[str], rows: list[list[str]]) -> None:
    from openpyxl.styles import Font

    bold = Font(bold=True)
    for col, value in enumerate(header, start=1):
        plain, _wants_bold = _cell_value(value)
        cell = ws.cell(row=1, column=col, value=plain)
        cell.font = bold  # header row is always bold
    for row_index, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            plain, wants_bold = _cell_value(value)
            cell = ws.cell(row=row_index, column=col, value=plain)
            if wants_bold:
                cell.font = bold
    # Light auto-width (capped) — measure stripped text.
    for col in range(1, len(header) + 1):
        letter = ws.cell(row=1, column=col).column_letter
        widest = len(str(_cell_value(header[col - 1] or "")[0]))
        for row in rows[:50]:
            if col - 1 < len(row):
                widest = max(widest, len(str(_cell_value(row[col - 1] or "")[0])))
        ws.column_dimensions[letter].width = min(max(widest + 2, 10), 48)


def build_xlsx(markdown: str) -> bytes:
    """Turn Markdown tables / CSV into an `.xlsx` byte payload."""
    from openpyxl import Workbook

    wb = Workbook()
    default = wb.active
    assert default is not None
    used_names: set[str] = set()

    def unique_title(raw: str | None, index: int) -> str:
        base = _sheet_title(raw, index)
        name = base
        n = 2
        while name.lower() in used_names:
            suffix = f" ({n})"
            name = (base[: 31 - len(suffix)] + suffix)[:31]
            n += 1
        used_names.add(name.lower())
        return name

    tables = _parse_markdown_tables(markdown)
    sheet_index = 0

    if tables:
        wb.remove(default)
        for heading, header, rows in tables:
            sheet_index += 1
            ws = wb.create_sheet(unique_title(heading, sheet_index))
            _write_sheet(ws, header, rows)
    elif _looks_like_csv(markdown):
        header, rows = _parse_csv(markdown)
        default.title = unique_title("Data", 1)
        _write_sheet(default, header, rows)
        sheet_index = 1
    else:
        from openpyxl.styles import Font

        default.title = unique_title("Notes", 1)
        default["A1"] = "Content"
        default["A1"].font = Font(bold=True)
        for row_index, line in enumerate(_prose_lines(markdown, skip_tables=False), start=2):
            plain, wants_bold = _cell_value(line)
            cell = default.cell(row=row_index, column=1, value=plain)
            if wants_bold:
                cell.font = Font(bold=True)
        default.column_dimensions["A"].width = 80
        sheet_index = 1

    # Prose outside tables → Notes sheet when we already built table sheets.
    if tables:
        notes = _prose_lines(markdown, skip_tables=True)
        if notes:
            from openpyxl.styles import Font

            ws = wb.create_sheet(unique_title("Notes", sheet_index + 1))
            ws["A1"] = "Notes"
            ws["A1"].font = Font(bold=True)
            for row_index, line in enumerate(notes, start=2):
                plain, wants_bold = _cell_value(line)
                cell = ws.cell(row=row_index, column=1, value=plain)
                if wants_bold:
                    cell.font = Font(bold=True)
            ws.column_dimensions["A"].width = 80

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
